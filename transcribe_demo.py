import whisper
from pydub import AudioSegment
from langchain.chat_models import ChatOpenAI
from langchain.prompts import PromptTemplate
from langchain.chains import LLMChain
import os
# for saving as docx
import markdown
from docx import Document
from bs4 import BeautifulSoup

from envVars.env_vars import openaiKey

# Step 1: Convert MP3 to WAV
def convert_mp3_to_wav(mp3_path, wav_path):
    audio = AudioSegment.from_mp3(mp3_path)
    audio.export(wav_path, format="wav")

# Step 2: Transcribe Audio using Whisper
def transcribe_audio(wav_path):
    model = whisper.load_model("base")  # or "medium"/"large" if you want better quality
    result = model.transcribe(wav_path)
    return result['text']

# Step 3: Use LangChain to generate dialogue from transcript
def generate_dialogue(transcript_text, llm):
    prompt = PromptTemplate(
        input_variables=["transcript"],
        template="make a dialog between nurse and patient from:\n\n{transcript}"
    )
    chain = LLMChain(llm=llm, prompt=prompt)
    return chain.run(transcript=transcript_text)

# Step 4: Use LangChain to convert dialogue to SOAP note
def generate_soap_note(dialogue_text, llm):
    prompt = PromptTemplate(
        input_variables=["dialogue"],
        template="Now make a doctors soap note from above, include subheadings:\n\n{dialogue}"
    )
    chain = LLMChain(llm=llm, prompt=prompt)
    return chain.run(dialogue=dialogue_text)

# Step 5: Main function to run the whole pipeline
def process_medical_audio(mp3_path, openai_api_key):
    wav_path = "temp.wav"
    os.environ["OPENAI_API_KEY"] = openai_api_key

    # Convert and transcribe
    convert_mp3_to_wav(mp3_path, wav_path)
    print("🔍 Transcribing...")
    transcript = transcribe_audio(wav_path)

    # Setup LLM
    llm = ChatOpenAI(temperature=0.7, model_name="gpt-4o")

    # Dialogue generation
    print("💬 Generating dialogue...")
    dialogue = generate_dialogue(transcript, llm)

    # SOAP note generation
    print("📝 Generating SOAP note...")
    soap_note = generate_soap_note(dialogue, llm)

    os.remove(wav_path)  # Cleanup
    return transcript, dialogue, soap_note

def process_multiple_audios():
    folder_path = "Files/Doctor Patient Interaction Audio Files"
    print(os.listdir(folder_path))
    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        if os.path.isfile(file_path):
            if ".DS" in file_path:
                continue

            transcript, dialogue, soap_note = process_medical_audio(file_path, api_key)

            mp3_folder = file_path.split("/")[-1].split(".")[0]
            output_folder_path = f"Files/Doctor Patient Outputs/{mp3_folder}"
            print (mp3_folder)

            os.makedirs(output_folder_path, exist_ok=True)
            writeTextToDocx(transcript, f"{output_folder_path}/transcription.docx")
            writeMarkdownToDocx(dialogue, f"{output_folder_path}/dialogue.docx")
            writeMarkdownToDocx(soap_note, f"{output_folder_path}/soap_note.docx")


def writeTextToDocx(text, file_path):
    # Create a new Word document
    doc = Document()
    doc.add_paragraph(text.strip())

    file_path = os.path.join(os.getcwd(), file_path)
    print(file_path)
    doc.save(file_path)

def writeMarkdownToDocx(md_text, file_path):
    html = markdown.markdown(md_text)

    # Step 2: Parse HTML and write to DOCX
    doc = Document()
    soup = BeautifulSoup(html, 'html.parser')

    for element in soup.children:
        if element.name == 'h1':
            doc.add_heading(element.text, level=1)
        elif element.name == 'h2':
            doc.add_heading(element.text, level=2)
        elif element.name == 'p':
            doc.add_paragraph(element.text)
        elif element.name == 'ul':
            for li in element.find_all('li'):
                doc.add_paragraph(li.text, style='List Bullet')

    file_path = os.path.join(os.getcwd(), file_path)
    # Save DOCX file
    doc.save(file_path)

# Example usage
if __name__ == "__main__":
    api_key = openaiKey
    process_multiple_audios()



    # transcript, dialogue, soap_note = process_medical_audio(mp3_file, api_key)

    # print("\n--- 🗒️ Transcript ---\n", transcript)
    # print("\n--- 🩺 Dialogue ---\n", dialogue)
    # print("\n--- 📋 SOAP Note ---\n", soap_note)
